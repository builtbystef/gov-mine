import io
from collections.abc import Callable

import docx
import openpyxl
import pptx

from gov_mine.integrations.parser import parse
from gov_mine.integrations.parser.files import DOCX, PPTX, XLSX


def test_files_are_parsed_page_by_page(make_pdf: Callable[..., bytes]):
    deck = pptx.Presentation()
    for title in ("Site plan", "Timeline"):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = title
    deck_bytes = io.BytesIO()
    deck.save(deck_bytes)
    book = openpyxl.Workbook()
    book.active.title = "Budget"
    book.active.append(["Item", "Cost"])
    book.create_sheet("Schedule").append(["Phase", "Year"])
    book_bytes = io.BytesIO()
    book.save(book_bytes)
    letter = docx.Document()
    letter.add_paragraph("Parkdale Hub community meeting")
    letter_bytes = io.BytesIO()
    letter.save(letter_bytes)

    pdf = parse(make_pdf("Parkdale Hub", "Second page"), "application/pdf")
    slides = parse(deck_bytes.getvalue(), PPTX)
    sheets = parse(book_bytes.getvalue(), XLSX)
    word = parse(letter_bytes.getvalue(), DOCX)
    csv = parse(b"lot,area\nA,2\n", "text/csv")
    image = parse(b"\x89PNG\r\n", "image/png")

    assert pdf is not None
    assert pdf.parser == "pdf:markitdown@1"
    assert [page.strip() for page in pdf.pages] == ["Parkdale Hub", "Second page"]
    assert slides is not None
    assert len(slides.pages) == 2
    assert "Site plan" in slides.pages[0]
    assert "Timeline" in slides.pages[1]
    assert sheets is not None
    assert len(sheets.pages) == 2
    assert sheets.pages[0].startswith("## Budget")
    assert sheets.pages[1].startswith("## Schedule")
    assert word is not None
    assert word.pages == ["Parkdale Hub community meeting\n"]
    assert csv is not None
    assert csv.pages == ["lot,area\nA,2\n"]
    assert image is None
